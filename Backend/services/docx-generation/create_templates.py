"""
Generate the standard .docx template files for each document type.

Run once to populate the templates/ directory:
    python create_templates.py

Each template uses Jinja2-style {{ TAG }} placeholders that docxtpl
will fill in at render time.
"""

import os
from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

TEMPLATE_DIR = Path(__file__).resolve().parent / "templates"
TEMPLATE_DIR.mkdir(exist_ok=True)


def _add_header(doc: Document, title: str) -> None:
    """Add a centered document title."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title)
    run.bold = True
    run.font.size = Pt(16)
    run.font.color.rgb = RGBColor(0x22, 0x22, 0x22)


def _add_line(doc: Document, label: str, value: str) -> None:
    """Add a single label: value line."""
    p = doc.add_paragraph()
    run_label = p.add_run(f"{label}: ")
    run_label.bold = True
    run_label.font.size = Pt(11)
    run_value = p.add_run(value)
    run_value.font.size = Pt(11)


def _add_spacer(doc: Document) -> None:
    doc.add_paragraph()


def _add_signature_block(doc: Document) -> None:
    """Add a signature line at the bottom."""
    _add_spacer(doc)
    _add_spacer(doc)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run("_________________________")
    run.font.size = Pt(11)

    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run2 = p2.add_run("Attending Physician")
    run2.font.size = Pt(10)
    run2.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    p3 = doc.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run3 = p3.add_run("License No.: {{ LICENSE_NO }}")
    run3.font.size = Pt(10)


def create_medical_certificate() -> None:
    doc = Document()
    _add_header(doc, "MEDICAL CERTIFICATE")
    _add_spacer(doc)
    _add_line(doc, "Date", "{{ DATE }}")
    _add_spacer(doc)
    p = doc.add_paragraph()
    run = p.add_run("This is to certify that ")
    run.font.size = Pt(11)
    run2 = p.add_run("{{ NAME }}")
    run2.bold = True
    run2.font.size = Pt(11)
    run3 = p.add_run(
        " has been examined and found to have the following diagnosis:"
    )
    run3.font.size = Pt(11)
    _add_spacer(doc)
    _add_line(doc, "Diagnosis", "{{ DIAGNOSIS }}")
    _add_signature_block(doc)
    doc.save(str(TEMPLATE_DIR / "medical_certificate.docx"))
    print("  Created: medical_certificate.docx")


def create_medical_clearance() -> None:
    doc = Document()
    _add_header(doc, "MEDICAL CLEARANCE")
    _add_spacer(doc)
    _add_line(doc, "Date", "{{ DATE }}")
    _add_spacer(doc)
    p = doc.add_paragraph()
    run = p.add_run("This is to certify that ")
    run.font.size = Pt(11)
    run2 = p.add_run("{{ NAME }}")
    run2.bold = True
    run2.font.size = Pt(11)
    run3 = p.add_run(
        " has been examined and found to be physically fit and cleared for the following purpose:"
    )
    run3.font.size = Pt(11)
    _add_spacer(doc)
    _add_line(doc, "Purpose", "{{ PURPOSE }}")
    _add_signature_block(doc)
    doc.save(str(TEMPLATE_DIR / "medical_clearance.docx"))
    print("  Created: medical_clearance.docx")


def create_prescription() -> None:
    doc = Document()
    _add_header(doc, "PRESCRIPTION")
    _add_spacer(doc)
    _add_line(doc, "Date", "{{ DATE }}")
    _add_line(doc, "Patient", "{{ NAME }}")
    _add_spacer(doc)

    p = doc.add_paragraph()
    run = p.add_run("Rx:")
    run.bold = True
    run.font.size = Pt(13)

    _add_spacer(doc)
    p2 = doc.add_paragraph()
    run2 = p2.add_run("{{ MEDICATION }}")
    run2.font.size = Pt(11)

    _add_signature_block(doc)

    # Extra line for PTR
    p_ptr = doc.add_paragraph()
    p_ptr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run_ptr = p_ptr.add_run("PTR No.: {{ PTR_NO }}")
    run_ptr.font.size = Pt(10)

    doc.save(str(TEMPLATE_DIR / "prescription.docx"))
    print("  Created: prescription.docx")


def create_lab_referral() -> None:
    doc = Document()
    _add_header(doc, "LABORATORY REFERRAL")
    _add_spacer(doc)
    _add_line(doc, "Date", "{{ DATE }}")
    _add_line(doc, "Patient", "{{ NAME }}")
    _add_spacer(doc)
    p = doc.add_paragraph()
    run = p.add_run(
        "The above-named patient is hereby referred for the following laboratory examination:"
    )
    run.font.size = Pt(11)
    _add_spacer(doc)
    _add_line(doc, "Test", "{{ TEST_TYPE }}")
    _add_signature_block(doc)
    doc.save(str(TEMPLATE_DIR / "lab_referral.docx"))
    print("  Created: lab_referral.docx")


if __name__ == "__main__":
    print(f"Creating templates in {TEMPLATE_DIR} ...")
    create_medical_certificate()
    create_medical_clearance()
    create_prescription()
    create_lab_referral()
    print("Done.")
